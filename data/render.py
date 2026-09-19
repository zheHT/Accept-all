"""Attachment renderers: SI and BL as .txt, .pdf, .docx, .xlsx.

The same shipment/BL dict is rendered with DIFFERENT field labels on the SI
vs the BL (drawn from pools.LABELS) so the extraction step must normalise
synonyms ("Load Port" == "Port of Loading").
"""
import random
import pools


def _lbl(rng, field):
    return rng.choice(pools.LABELS[field])


# ---------------------------------------------------------------------------
# Plain text
# ---------------------------------------------------------------------------
def si_txt(rng, ship):
    r = ship["refs"]
    lines = [
        "SHIPPING INSTRUCTION",
        "=" * 40,
        "",
        f"{_lbl(rng,'shipper')}: {ship['shipper']}",
        f"  {'; '.join(ship['shipper_addr'])}",
        f"{_lbl(rng,'consignee')}: {ship['consignee']}",
        f"  {'; '.join(ship['consignee_addr'])}",
        f"{_lbl(rng,'notify_party')}: {ship['notify_party']}",
        f"{_lbl(rng,'port_of_loading')}: {ship['port_of_loading']} ({ship['pol_code']})",
        f"{_lbl(rng,'port_of_discharge')}: {ship['port_of_discharge']} ({ship['pod_code']})",
        f"{_lbl(rng,'container_count')}: {ship['container_count']} x {ship['container_size']}",
        f"{_lbl(rng,'gross_weight_kg')}: {ship['gross_weight_kg']:,} KG",
        f"{_lbl(rng,'vessel')}: {ship['vessel']}",
        f"{_lbl(rng,'voyage')}: {ship['voyage']}",
        f"{_lbl(rng,'commodity')}: {ship['commodity']}",
        f"HS Code: {ship['hs_code']}",
        f"{_lbl(rng,'booking')}: {r['booking']}",
        f"OC No.: {r['oc_ref']}",
        "Freight: PREPAID",
    ]
    return "\n".join(lines) + "\n"


def bl_txt(rng, bl):
    r = bl["refs"]
    lines = [
        "BILL OF LADING (DRAFT)",
        "=" * 40,
        "",
        f"{_lbl(rng,'shipper')}: {bl['shipper']}",
        f"  {'; '.join(bl['shipper_addr'])}",
        f"{_lbl(rng,'consignee')}: {bl['consignee']}",
        f"  {'; '.join(bl['consignee_addr'])}",
        f"{_lbl(rng,'notify_party')}: {bl['notify_party']}",
        f"{_lbl(rng,'port_of_loading')}: {bl['port_of_loading']} ({bl['pol_code']})",
        f"{_lbl(rng,'port_of_discharge')}: {bl['port_of_discharge']} ({bl['pod_code']})",
        f"{_lbl(rng,'container_count')}: {bl['container_count']} x {bl['container_size']}",
        f"{_lbl(rng,'gross_weight_kg')}: {bl['gross_weight_kg']:,} KG",
        f"{_lbl(rng,'vessel')}: {bl['vessel']}",
        f"{_lbl(rng,'voyage')}: {bl['voyage']}",
        f"{_lbl(rng,'commodity')}: {bl['commodity']}",
        f"{_lbl(rng,'bl_no')}: {r['bl_no']}",
        f"{_lbl(rng,'booking')}: {r['booking']}",
        "Freight: PREPAID",
    ]
    return "\n".join(lines) + "\n"


def write_txt(path, text):
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)


# ---------------------------------------------------------------------------
# PDF  (mirrors the real BDP-style BL / MSC BL-instruction layout)
# ---------------------------------------------------------------------------
def write_pdf(path, doc, title, rng):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(path, pagesize=A4)
    w, h = A4
    y = h - 20 * mm
    c.setFont("Helvetica-Bold", 12)
    c.drawString(20 * mm, y, title)
    y -= 6 * mm
    c.setFont("Helvetica", 8)
    c.drawString(20 * mm, y, f"B/L NUMBER: {doc['refs']['bl_no']}    BOOKING NO. {doc['refs']['booking']}")
    y -= 8 * mm

    def block(label, value_lines):
        nonlocal y
        c.setFont("Helvetica-Bold", 8)
        c.drawString(20 * mm, y, label)
        c.setFont("Helvetica", 8)
        for ln in value_lines:
            c.drawString(60 * mm, y, ln[:70])
            y -= 4.5 * mm
        y -= 1.5 * mm

    block(_lbl(rng, "shipper"), [doc["shipper"]] + doc["shipper_addr"])
    block(_lbl(rng, "consignee"), [doc["consignee"]] + doc["consignee_addr"])
    block(_lbl(rng, "notify_party"), [doc["notify_party"]] + doc["notify_addr"])
    block(_lbl(rng, "port_of_loading"), [f"{doc['port_of_loading']}"])
    block(_lbl(rng, "port_of_discharge"), [f"{doc['port_of_discharge']}"])
    block(_lbl(rng, "vessel"), [f"{doc['vessel']}"])
    y -= 4 * mm

    # container table
    c.setFont("Helvetica-Bold", 8)
    c.drawString(20 * mm, y, "CONTAINER NO.")
    c.drawString(70 * mm, y, "DESCRIPTION")
    c.drawString(140 * mm, y, "GROSS WEIGHT (KG)")
    y -= 5 * mm
    c.setFont("Helvetica", 8)
    per = doc["gross_weight_kg"] // doc["container_count"]
    for i in range(doc["container_count"]):
        cn = "".join(random.choice("ABCDEFGHIJKLMNOPQRSTUVWXYZ") for _ in range(4)) + "".join(random.choice("0123456789") for _ in range(7))
        wt = per if i < doc["container_count"] - 1 else doc["gross_weight_kg"] - per * (doc["container_count"] - 1)
        c.drawString(20 * mm, y, cn)
        c.drawString(70 * mm, y, f"{doc['container_size']} {doc['commodity'][:30]}")
        c.drawString(140 * mm, y, f"{wt:,}")
        y -= 4.5 * mm
        if y < 30 * mm:
            break
    y -= 3 * mm
    c.setFont("Helvetica-Bold", 8)
    c.drawString(20 * mm, y, f"{_lbl(rng,'container_count')}: {doc['container_count']} x {doc['container_size']}")
    y -= 5 * mm
    c.drawString(20 * mm, y, f"TOTAL {_lbl(rng,'gross_weight_kg')}: {doc['gross_weight_kg']:,} KG")
    y -= 5 * mm
    c.setFont("Helvetica", 7)
    c.drawString(20 * mm, y, f"HS CODE {doc['hs_code']}   FREIGHT PREPAID   OC NO. {doc['refs']['oc_ref']}")
    c.showPage()
    c.save()


# ---------------------------------------------------------------------------
# DOCX  (mirrors the Chinese/English bilingual BL word doc)
# ---------------------------------------------------------------------------
def write_docx(path, doc, rng):
    from docx import Document
    d = Document()
    d.add_heading("BILL OF LADING (DRAFT)", level=1)
    d.add_paragraph(f"B/L NO.(提单号): {doc['refs']['bl_no']}")
    t = d.add_table(rows=0, cols=2)
    t.style = "Table Grid"

    def row(label, value):
        cells = t.add_row().cells
        cells[0].text = label
        cells[1].text = value

    row(f"{_lbl(rng,'shipper')} (发货人)", doc["shipper"] + "\n" + "\n".join(doc["shipper_addr"]))
    row(f"{_lbl(rng,'consignee')} (收货人)", doc["consignee"] + "\n" + "\n".join(doc["consignee_addr"]))
    row(f"{_lbl(rng,'notify_party')} (通知人)", doc["notify_party"] + "\n" + "\n".join(doc["notify_addr"]))
    row(f"{_lbl(rng,'port_of_loading')} (装货港)", doc["port_of_loading"])
    row(f"{_lbl(rng,'port_of_discharge')} (卸货港)", doc["port_of_discharge"])
    row(f"{_lbl(rng,'container_count')} (箱数)", f"{doc['container_count']} x {doc['container_size']}")
    row(f"{_lbl(rng,'gross_weight_kg')} (毛重 KGS)", f"{doc['gross_weight_kg']:,}")
    row(f"{_lbl(rng,'vessel')} (船名)", f"{doc['vessel']}")
    row(f"{_lbl(rng,'commodity')} (货名)", doc["commodity"])
    d.add_paragraph(f"ORDER NO.: {doc['refs']['so_number']}   FREIGHT PREPAID")
    d.save(path)


# ---------------------------------------------------------------------------
# XLSX  (mirrors the SI / BL excel worksheet)
# ---------------------------------------------------------------------------
def write_xlsx(path, doc, kind, rng):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S.I." if kind == "SI" else "BL"
    rows = [
        (doc["shipper"], ""),
        ("", ""),
        ("BL INSTRUCTION" if kind == "SI" else "BILL OF LADING", doc["refs"]["so_number"]),
        (_lbl(rng, "shipper"), doc["shipper"] + " | " + "; ".join(doc["shipper_addr"])),
        (_lbl(rng, "consignee"), doc["consignee"] + " | " + "; ".join(doc["consignee_addr"])),
        (_lbl(rng, "notify_party"), doc["notify_party"] + " | " + "; ".join(doc["notify_addr"])),
        (_lbl(rng, "port_of_loading"), doc["port_of_loading"]),
        (_lbl(rng, "port_of_discharge"), doc["port_of_discharge"]),
        (_lbl(rng, "container_count"), f"{doc['container_count']} x {doc['container_size']}"),
        (_lbl(rng, "gross_weight_kg"), doc["gross_weight_kg"]),
        (_lbl(rng, "vessel"), doc["vessel"]),
        (_lbl(rng, "commodity"), doc["commodity"]),
        ("HS CODE", doc["hs_code"]),
        (_lbl(rng, "bl_no") if kind == "BL" else _lbl(rng, "booking"),
         doc["refs"]["bl_no"] if kind == "BL" else doc["refs"]["booking"]),
        ("FREIGHT", "PREPAID"),
    ]
    for label, value in rows:
        ws.append([label, value])
    wb.save(path)
