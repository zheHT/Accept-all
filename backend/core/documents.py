from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".txt"}


def extract_text(filename: str, data: bytes) -> tuple[str, bool]:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        return "", False
    try:
        if suffix == ".txt":
            text = data.decode("utf-8", errors="replace")
        elif suffix == ".pdf":
            reader = PdfReader(BytesIO(data))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx":
            document = Document(BytesIO(data))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            tables = [
                " | ".join(cell.text for cell in row.cells)
                for table in document.tables
                for row in table.rows
            ]
            text = "\n".join([*paragraphs, *tables])
        else:
            workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
            rows = []
            for sheet in workbook.worksheets:
                rows.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    rows.append(" | ".join("" if cell is None else str(cell) for cell in row))
            text = "\n".join(rows)
    except Exception:
        return "", False
    text = text.strip()
    return text, bool(text)
