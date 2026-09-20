"""Exhaustive automated unit test suite for Maritime Shipping Email Classifier & Triage Agent.

Uses pytest, unittest.mock, and httpx / TestClient.
All external LLM calls are mocked to ensure deterministic, network-free execution
while strictly validating production model configuration.
"""
import io
import json
import sys
from pathlib import Path
from unittest.mock import MagicMock, patch

# Ensure project root is in sys.path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))


import docx
import openpyxl
import pytest
from fastapi.testclient import TestClient
from pydantic import ValidationError

from backend.evaluation_adapter.attachment_sniffer import MAX_CHARS, inspect_attachment
from backend.evaluation_adapter.classifier_flow import (
    _check_missing_attachments,
    classify_email,
)
from backend.evaluation_adapter.main import app
from backend.evaluation_adapter.schemas import (
    EmailCategory,
    EmailClassification,
)


# --------------------------------------------------------------------------
# Fixtures
# --------------------------------------------------------------------------
@pytest.fixture
def mock_gemini_client():
    """Mock fixture for google.genai Client to prevent external network calls."""
    with patch("backend.evaluation_adapter.classifier_flow.genai.Client") as mock_client_cls, \
         patch.dict("os.environ", {"GEMINI_API_KEY": "AIzaSy_mock_api_key_test_12345"}):
        mock_client_instance = MagicMock()
        mock_client_cls.return_value = mock_client_instance

        default_resp = MagicMock()
        default_resp.text = json.dumps({
            "category": "DOCUMENT_COMPARISON",
            "confidence": 0.98,
            "reasoning": "Email contains draft BL and SI for cross-check.",
            "is_comparison_candidate": True,
            "missing_attachments_flag": False,
            "detected_attachments": ["draft_bl.pdf", "si.txt"],
        })
        mock_client_instance.models.generate_content.return_value = default_resp
        yield mock_client_instance


@pytest.fixture
def api_client():
    """HTTP test client using Starlette / httpx."""
    return TestClient(app)


# --------------------------------------------------------------------------
# Area 1: Pydantic Schema & Bounds Validation
# --------------------------------------------------------------------------
def test_schema_valid_categories():
    """1. Verify EmailClassification accepts all 5 operational categories with valid fields."""
    expected_categories = [
        EmailCategory.DOCUMENT_COMPARISON,
        EmailCategory.NEW_SI_REQUEST,
        EmailCategory.INVOICE_QUERY,
        EmailCategory.GENERAL,
        EmailCategory.SPAM,
    ]
    assert len(expected_categories) == 5

    for cat in expected_categories:
        item = EmailClassification(
            category=cat,
            confidence=0.92,
            reasoning=f"Valid classification for {cat.value}",
            is_comparison_candidate=(cat == EmailCategory.DOCUMENT_COMPARISON),
            missing_attachments_flag=False,
            detected_attachments=["doc1.txt"],
        )
        assert item.category == cat
        assert item.category.value == cat.value
        assert isinstance(item.confidence, float)
        assert item.is_comparison_candidate == (cat == EmailCategory.DOCUMENT_COMPARISON)


def test_schema_confidence_out_of_bounds():
    """2. Assert Pydantic raises ValidationError for confidence > 1.0 or < 0.0."""
    with pytest.raises(ValidationError):
        EmailClassification(
            category=EmailCategory.GENERAL,
            confidence=1.25,  # Out of bounds (> 1.0)
            reasoning="Confidence above upper bound",
            is_comparison_candidate=False,
        )

    with pytest.raises(ValidationError):
        EmailClassification(
            category=EmailCategory.GENERAL,
            confidence=-0.10,  # Out of bounds (< 0.0)
            reasoning="Confidence below lower bound",
            is_comparison_candidate=False,
        )


def test_schema_attribute_defaults():
    """3. Verify optional attribute defaults: missing_attachments_flag=False, detected_attachments=[]."""
    item = EmailClassification(
        category=EmailCategory.NEW_SI_REQUEST,
        confidence=0.88,
        reasoning="New SI submission without carrier draft BL.",
        is_comparison_candidate=False,
    )
    assert item.missing_attachments_flag is False
    assert item.detected_attachments == []
    assert isinstance(item.detected_attachments, list)


# --------------------------------------------------------------------------
# Area 2: Model Configuration Verification
# --------------------------------------------------------------------------
def test_agent_uses_configured_gemini_model(mock_gemini_client, monkeypatch):
    """The classifier passes the shared model setting to the SDK."""
    monkeypatch.setenv("GEMINI_MODEL", "gemini-test-flash")
    # Execute classification call
    classify_email(
        email_id="model_cfg_test",
        subject="Draft BL and SI cross check",
        sender="ops@ocean.com",
        body="Please verify the draft BL against our SI.",
        attachment_previews={"draft_bl.txt": "Draft content", "si.txt": "SI content"},
    )

    # Verify generate_content was called
    assert mock_gemini_client.models.generate_content.called
    call_args, call_kwargs = mock_gemini_client.models.generate_content.call_args

    # Assert model parameter comes from the shared environment setting.
    model_passed = call_kwargs.get("model")
    assert model_passed == "gemini-test-flash"

    # Assert deterministic configuration (temperature=0.0, response_schema=EmailClassification)
    config = call_kwargs.get("config")
    assert config is not None
    assert config.temperature == 0.0
    assert config.response_mime_type == "application/json"
    assert config.response_schema == EmailClassification
    assert mock_gemini_client.close.called


def test_classifier_closes_client_after_sdk_failure(mock_gemini_client):
    """A failed live classification still releases its per-call SDK client."""
    mock_gemini_client.models.generate_content.side_effect = RuntimeError("sdk failed")

    with pytest.raises(RuntimeError, match="sdk failed"):
        classify_email(
            email_id="client_cleanup_failure",
            subject="Draft BL check",
            sender="ops@example.com",
            body="Please compare the draft BL and SI.",
            attachment_previews={},
        )

    assert mock_gemini_client.close.called


def test_classifier_preserves_request_error_when_client_close_fails(mock_gemini_client):
    """A cleanup failure cannot replace the Gemini request failure."""
    mock_gemini_client.models.generate_content.side_effect = RuntimeError("request failed")
    mock_gemini_client.close.side_effect = RuntimeError("close failed")

    with pytest.raises(RuntimeError, match="request failed"):
        classify_email(
            email_id="client_cleanup_request_failure",
            subject="Draft BL check",
            sender="ops@example.com",
            body="Please compare the draft BL and SI.",
            attachment_previews={},
        )


def test_classifier_returns_result_when_client_close_fails(mock_gemini_client):
    """A cleanup failure after a valid response does not discard classification."""
    mock_gemini_client.close.side_effect = RuntimeError("close failed")

    result = classify_email(
        email_id="client_cleanup_success",
        subject="Draft BL check",
        sender="ops@example.com",
        body="Please compare the draft BL and SI.",
        attachment_previews={
            "si.pdf": "[PDF document: original bytes retained]",
            "bl.pdf": "[PDF document: original bytes retained]",
        },
    )

    assert result.category == EmailCategory.DOCUMENT_COMPARISON


# --------------------------------------------------------------------------
# Area 3: Attachment Sniffer Resilience
# --------------------------------------------------------------------------
def test_sniffer_text_truncation_ceiling():
    """5. Verify text previews are cleanly truncated to the MAX_CHARS ceiling (1,000 characters)."""
    long_payload = "Port of Loading: Singapore | Container: MSCU1234567 | " * 60
    assert len(long_payload) > 1500

    preview = inspect_attachment("manifest.txt", long_payload)
    assert len(preview) <= MAX_CHARS
    assert len(preview) == 1000
    assert preview.startswith("Port of Loading: Singapore")


def test_sniffer_corrupted_binary_handling():
    """6. Ensure corrupted/non-decodable binary junk returns safe indicator without crashing."""
    corrupted_bytes = b"\x00\xff\xfe\xca\xfe\xba\xbe\xef\x00\x01\x02\x03"
    preview = inspect_attachment("broken_doc.pdf", corrupted_bytes)
    assert preview in [
        "[Unreadable or Corrupted File]",
        "[Scanned or Image-only PDF]",
    ]


def test_sniffer_supported_extensions_preview():
    """7. Verify valid mock payloads for .docx, .xlsx, and text generate valid previews."""
    # 1. Clean plain text
    txt_preview = inspect_attachment("si.txt", "Shipper: Moorim Paper\nConsignee: PaperOne\n")
    assert "Shipper: Moorim Paper" in txt_preview

    # 2. In-memory Word document (.docx)
    doc = docx.Document()
    doc.add_paragraph("Draft Bill of Lading #MEDU104332")
    doc.add_paragraph("Shipper: April Fine Paper Trading")
    doc.add_paragraph("Port of Discharge: Callao, Peru")
    docx_buf = io.BytesIO()
    doc.save(docx_buf)
    docx_preview = inspect_attachment("draft_bl.docx", docx_buf.getvalue())
    assert "Draft Bill of Lading" in docx_preview
    assert "April Fine Paper" in docx_preview

    # 3. In-memory Excel spreadsheet (.xlsx)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ShippingInstructions"
    ws.append(["Shipper", "Consignee", "Port of Loading", "Gross Weight"])
    ws.append(["Acme Logistics", "Pacific Global", "Singapore", "24,000 KG"])
    xlsx_buf = io.BytesIO()
    wb.save(xlsx_buf)
    xlsx_preview = inspect_attachment("si_sheet.xlsx", xlsx_buf.getvalue())
    assert "ShippingInstructions" in xlsx_preview
    assert "Acme Logistics" in xlsx_preview


# --------------------------------------------------------------------------
# Area 4: Cognitive Routing & Guardrail Logic
# --------------------------------------------------------------------------
def test_guardrail_document_comparison_missing_attachments(mock_gemini_client):
    """8. DOCUMENT_COMPARISON intent with no attachments must set missing_attachments_flag=True."""
    mock_resp = MagicMock()
    mock_resp.text = json.dumps({
        "category": "DOCUMENT_COMPARISON",
        "confidence": 0.95,
        "reasoning": "Sender requests checking draft BL against SI, but no files were attached.",
        "is_comparison_candidate": True,
        "missing_attachments_flag": False,  # LLM initially emits False
        "detected_attachments": [],
    })
    mock_gemini_client.models.generate_content.return_value = mock_resp

    result = classify_email(
        email_id="missing_att_01",
        subject="Please check draft BL against SI",
        sender="importer@tradex.com",
        body="Dear customer service, please check draft BL against SI. Forgot to attach.",
        attachment_previews={},  # Empty attachments
    )

    assert result.category == EmailCategory.DOCUMENT_COMPARISON
    assert result.is_comparison_candidate is True
    # Guardrail must force missing_attachments_flag to True
    assert result.missing_attachments_flag is True


def test_guardrail_document_comparison_with_valid_attachments(mock_gemini_client):
    """9. DOCUMENT_COMPARISON intent with valid readable attachments must set missing_attachments_flag=False."""
    mock_resp = MagicMock()
    mock_resp.text = json.dumps({
        "category": "DOCUMENT_COMPARISON",
        "confidence": 0.97,
        "reasoning": "Both carrier draft BL and customer SI are present and readable.",
        "is_comparison_candidate": True,
        "missing_attachments_flag": False,
        "detected_attachments": ["SI_draft.pdf", "Carrier_BL.txt"],
    })
    mock_gemini_client.models.generate_content.return_value = mock_resp

    result = classify_email(
        email_id="valid_att_01",
        subject="TO CONFIRM DOCS _ 5RSG-00133",
        sender="ops@ocean.com",
        body="Attached are the SI and draft BL for OC 5RSG-00133. Please check details.",
        attachment_previews={
            "SI_draft.pdf": "Shipper: Moorim Paper | Consignee: PaperOne Inc",
            "Carrier_BL.txt": "Carrier Draft Bill of Lading | MEDU104332",
        },
    )

    assert result.category == EmailCategory.DOCUMENT_COMPARISON
    assert result.is_comparison_candidate is True
    assert result.missing_attachments_flag is False
    assert len(result.detected_attachments) == 2


def test_guardrail_clears_llm_missing_flag_for_two_pdf_markers(mock_gemini_client):
    """Two previews override an incorrect LLM missing-attachment flag."""
    mock_response = MagicMock()
    mock_response.text = json.dumps({
        "category": "DOCUMENT_COMPARISON",
        "confidence": 0.97,
        "reasoning": "Both documents were supplied for downstream review.",
        "is_comparison_candidate": True,
        "missing_attachments_flag": True,
        "detected_attachments": ["si.pdf", "bl.pdf"],
    })
    mock_gemini_client.models.generate_content.return_value = mock_response

    result = classify_email(
        email_id="two_pdf_markers",
        subject="Review draft BL",
        sender="ops@example.com",
        body="Please compare the draft BL with our SI.",
        attachment_previews={
            "si.pdf": "[PDF document: original bytes retained]",
            "bl.pdf": "[PDF document: original bytes retained]",
        },
    )

    assert result.missing_attachments_flag is False


def test_scanned_pdf_markers_are_not_missing_attachments():
    """Two supplied document previews count even when their PDFs need review."""
    assert _check_missing_attachments(
        EmailCategory.DOCUMENT_COMPARISON,
        {
            "si.pdf": "[PDF document: original bytes retained]",
            "bl.pdf": "[PDF document: original bytes retained]",
        },
    ) is False


def test_one_expected_attachment_is_missing():
    """A comparison has a missing attachment when fewer than two previews arrive."""
    assert _check_missing_attachments(
        EmailCategory.DOCUMENT_COMPARISON,
        {"si.txt": "SHIPPING INSTRUCTION"},
    ) is True


def test_deceptive_subject_line_resolution(mock_gemini_client):
    """10. Body intent must override deceptive subject lines (e.g. 'General Inquiry' -> DOCUMENT_COMPARISON)."""
    mock_resp = MagicMock()
    mock_resp.text = json.dumps({
        "category": "DOCUMENT_COMPARISON",
        "confidence": 0.94,
        "reasoning": "Despite general inquiry subject, body explicitly asks to cross-check draft BL vs SI.",
        "is_comparison_candidate": True,
        "missing_attachments_flag": False,
        "detected_attachments": ["draft_bl.pdf", "si.docx"],
    })
    mock_gemini_client.models.generate_content.return_value = mock_resp

    result = classify_email(
        email_id="deceptive_subj_01",
        subject="General Inquiry regarding booking",
        sender="shipper@client.com",
        body="Please cross-check our SI against carrier draft BL and verify container weights.",
        attachment_previews={
            "draft_bl.pdf": "Draft BL MEDU99210",
            "si.docx": "Shipping Instruction BKG889",
        },
    )

    assert result.category == EmailCategory.DOCUMENT_COMPARISON
    assert result.is_comparison_candidate is True
    assert "cross-check" in result.reasoning.lower() or "draft bl" in result.reasoning.lower()


def test_new_si_submission_without_bl(mock_gemini_client):
    """11. Initial Shipping Instructions filing without draft BL must classify as NEW_SI_REQUEST."""
    mock_resp = MagicMock()
    mock_resp.text = json.dumps({
        "category": "NEW_SI_REQUEST",
        "confidence": 0.96,
        "reasoning": "Customer submitting initial SI for booking creation without any carrier draft BL.",
        "is_comparison_candidate": False,
        "missing_attachments_flag": False,
        "detected_attachments": ["SI_booking_123.xlsx"],
    })
    mock_gemini_client.models.generate_content.return_value = mock_resp

    result = classify_email(
        email_id="new_si_01",
        subject="New Shipping Instruction - Booking BKG-992100",
        sender="exporter@trade.com",
        body="Dear Carrier Ops, please find attached our initial Shipping Instruction for filing.",
        attachment_previews={"SI_booking_123.xlsx": "Sheets: SI | Shipper: Apex Cargo"},
    )

    assert result.category == EmailCategory.NEW_SI_REQUEST
    assert result.is_comparison_candidate is False
    assert result.missing_attachments_flag is False


# --------------------------------------------------------------------------
# Area 5: FastAPI HTTP Endpoints
# --------------------------------------------------------------------------
def test_api_health_endpoint(api_client):
    """12. GET /health endpoint returns HTTP 200 OK and service metadata."""
    response = api_client.get("/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "ok"
    assert data["service"] == "email_classifier_triage_agent"
    assert "gemini_client_ready" in data


def test_api_process_endpoint_e2e(api_client, mock_gemini_client):
    """13. POST /api/process (or /api/classify) responds with 200 OK and conforms to EmailClassification."""
    payload = {
        "email_id": "api_e2e_01",
        "subject": "TO CONFIRM DOCS _ 5RSG-00133",
        "body": "Attached are the SI and draft BL for OC 5RSG-00133. Please check details.",
        "sender": "willy@shipping.com",
        "attachments": ["email_001_SI.txt", "email_001_BL.txt"],
    }

    # Test /api/classify
    res1 = api_client.post("/api/classify", json=payload)
    assert res1.status_code == 200
    data1 = res1.json()
    assert data1["category"] == "DOCUMENT_COMPARISON"
    assert data1["is_comparison_candidate"] is True
    assert 0.0 <= data1["confidence"] <= 1.0
    assert isinstance(data1["reasoning"], str)

    # Test /api/process alias
    res2 = api_client.post("/api/process", json=payload)
    assert res2.status_code == 200
    data2 = res2.json()
    assert data2["category"] == "DOCUMENT_COMPARISON"
    assert data2["is_comparison_candidate"] is True


def test_api_invoice_and_spam_routing(mock_gemini_client):
    """14. Additional edge cases: Verify INVOICE_QUERY and SPAM cognitive classification."""
    # Invoice dispute
    mock_resp_inv = MagicMock()
    mock_resp_inv.text = json.dumps({
        "category": "INVOICE_QUERY",
        "confidence": 0.95,
        "reasoning": "Dispute of demurrage charges on ocean freight invoice.",
        "is_comparison_candidate": False,
        "missing_attachments_flag": False,
        "detected_attachments": ["INV-990.pdf"],
    })
    mock_gemini_client.models.generate_content.return_value = mock_resp_inv

    inv_result = classify_email(
        email_id="inv_01",
        subject="Dispute Demurrage - Container MSCU7721",
        sender="billing@importer.com",
        body="Please review freight billing invoice INV-990 and issue credit note.",
        attachment_previews={"INV-990.pdf": "Invoice INV-990 Demurrage $450"},
    )
    assert inv_result.category == EmailCategory.INVOICE_QUERY
    assert inv_result.is_comparison_candidate is False

    # Marketing Spam
    mock_resp_spam = MagicMock()
    mock_resp_spam.text = json.dumps({
        "category": "SPAM",
        "confidence": 0.99,
        "reasoning": "Unsolicited SEO marketing campaign.",
        "is_comparison_candidate": False,
        "missing_attachments_flag": False,
        "detected_attachments": [],
    })
    mock_gemini_client.models.generate_content.return_value = mock_resp_spam

    spam_result = classify_email(
        email_id="spam_01",
        subject="Get #1 Google Ranking with our SEO marketing",
        sender="promo@external.org",
        body="Special discount on SEO services. Unsubscribe to opt out.",
        attachment_previews={},
    )
    assert spam_result.category == EmailCategory.SPAM
    assert spam_result.is_comparison_candidate is False
