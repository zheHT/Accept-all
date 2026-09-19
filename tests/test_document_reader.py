from io import BytesIO

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from backend.extraction.document_reader import (
    UnreadableDocumentError,
    detect_document_kind,
    prepare_document,
)
from backend.extraction.schemas import DocumentInput, DocumentKind, PreparedDocument


FIELD_LABELS = (
    "Shipper",
    "Consignee",
    "Notify Party",
    "Port of Loading",
    "Port of Discharge",
    "Container Count",
    "Gross Weight",
)


def _field_lines() -> list[str]:
    return [
        "Shipper: APRIL FAR EAST",
        "Consignee: MOORIM SP",
        "Notify Party: UAB NOVAKOPA",
        "Port of Loading: PORT KLANG",
        "Port of Discharge: CALLAO",
        "Container Count: 1",
        "Gross Weight: 21577 KG",
    ]


def docx_bytes() -> bytes:
    stream = BytesIO()
    document = Document()
    document.add_heading("BILL OF LADING (DRAFT)", level=1)
    for line in _field_lines():
        document.add_paragraph(line)
    document.save(stream)
    return stream.getvalue()


def xlsx_bytes() -> bytes:
    stream = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["SHIPPING INSTRUCTION"])
    for line in _field_lines():
        label, value = line.split(": ", maxsplit=1)
        sheet.append([label, value])
    workbook.save(stream)
    return stream.getvalue()


def blank_pdf_bytes() -> bytes:
    stream = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.write(stream)
    return stream.getvalue()


@pytest.fixture
def valid_pdf_bytes() -> bytes:
    return blank_pdf_bytes()


@pytest.fixture
def image_only_pdf_bytes() -> bytes:
    return blank_pdf_bytes()


def test_txt_is_decoded_to_text():
    prepared = prepare_document(
        DocumentInput("si.txt", b"SHIPPING INSTRUCTION\nShipper: A")
    )

    assert prepared.text.startswith("SHIPPING INSTRUCTION")
    assert prepared.data is None


def test_valid_pdf_keeps_original_bytes(valid_pdf_bytes):
    prepared = prepare_document(DocumentInput("si.pdf", valid_pdf_bytes))

    assert prepared.mime_type == "application/pdf"
    assert prepared.data == valid_pdf_bytes
    assert prepared.text is None


def test_image_only_pdf_is_not_rejected(image_only_pdf_bytes):
    prepared = prepare_document(DocumentInput("scan.pdf", image_only_pdf_bytes))

    assert prepared.data == image_only_pdf_bytes


def test_corrupt_pdf_is_unreadable():
    with pytest.raises(UnreadableDocumentError):
        prepare_document(DocumentInput("broken.pdf", b"%PDF corrupt"))


@pytest.mark.parametrize(
    ("filename", "document_bytes", "title"),
    [
        ("draft_bl.docx", docx_bytes, "BILL OF LADING (DRAFT)"),
        ("shipping_instruction.xlsx", xlsx_bytes, "SHIPPING INSTRUCTION"),
    ],
)
def test_office_documents_are_converted_to_markdown_with_shipping_fields(
    filename, document_bytes, title
):
    prepared = prepare_document(DocumentInput(filename, document_bytes()))

    assert title in prepared.text
    assert all(label in prepared.text for label in FIELD_LABELS)
    assert prepared.data is None


@pytest.mark.parametrize("filename", ["broken.docx", "broken.xlsx"])
def test_corrupt_office_document_is_unreadable(filename):
    with pytest.raises(UnreadableDocumentError):
        prepare_document(DocumentInput(filename, b"not an Office container"))


def test_wrong_document_title_is_detected():
    prepared = PreparedDocument(
        "claimed_bl.txt", "text/plain", text="COMMERCIAL INVOICE\nInvoice No: 1"
    )

    assert detect_document_kind(prepared) == DocumentKind.OTHER


def test_incidental_invoice_reference_does_not_reject_si():
    prepared = PreparedDocument(
        "si.txt",
        "text/plain",
        text="SHIPPING INSTRUCTION\nDocuments required: COMMERCIAL INVOICE",
    )

    assert detect_document_kind(prepared) == DocumentKind.SI


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("SHIPPING INSTRUCTION\nShipper: A", DocumentKind.SI),
        ("BILL OF LADING (DRAFT)\nShipper: A", DocumentKind.BL),
        ("Notes for shipping instruction\nShipper: A", DocumentKind.UNKNOWN),
    ],
)
def test_document_kind_uses_only_the_first_non_empty_line(text, expected):
    assert detect_document_kind(PreparedDocument("document.txt", "text/plain", text=text)) == expected


def test_pdf_document_kind_is_unknown():
    assert (
        detect_document_kind(PreparedDocument("si.pdf", "application/pdf", data=b"pdf"))
        == DocumentKind.UNKNOWN
    )
